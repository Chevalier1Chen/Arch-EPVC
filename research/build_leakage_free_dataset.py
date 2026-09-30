from __future__ import annotations

import json
import os
import shutil
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(r"C:\Users\DELL\Documents\论文")
SCI8 = Path(r"C:\Users\DELL\Desktop\SCI8")
SOURCE_XLSX = SCI8 / "数据集完成" / "numerical data.xlsx"
SOURCE_TS = SCI8 / "数据集完成" / "Time series data"
SOURCE_TS_FALLBACK = SCI8 / "Time series data"
SOURCE_VOXELS = ROOT / "external_validation_150" / "model_run_restored_inputs" / "training_voxels_32.npy"
SOURCE_BUILDINGS = ROOT / "school-building-design-platform" / "public" / "data" / "buildings.json"
SOURCE_GH = SCI8 / "shandong_teaching_campus_overview_package" / "data" / "gh_building_energy_inputs.csv"
SOURCE_OBJ = SCI8 / "shandong_teaching_campus_overview_package" / "models" / "per_campus_obj"
SOURCE_YEAR = SCI8 / "Simulation" / "gh_building_simulation_min_inputs.xlsx"
OVERLAP = ROOT / "external_validation_61_final" / "external_vs_training_coordinate_overlap_check.csv"

OUT = ROOT / "leakage_free_dataset_excluding_61_campuses"


def campus_id(building_id: str) -> str:
    return str(building_id).split("_T", 1)[0]


def hardlink_or_copy(source: Path, destination: Path) -> str:
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        return "existing"
    try:
        os.link(source, destination)
        return "hardlink"
    except OSError:
        shutil.copy2(source, destination)
        return "copy"


def filter_workbook(keep_ids: set[str], excluded_campuses: set[str]) -> None:
    destination = OUT / "source-data" / "numerical data.xlsx"
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(SOURCE_XLSX)
    sheet = workbook["selected_variables"]
    header = {sheet.cell(1, column).value: column for column in range(1, sheet.max_column + 1)}
    number_column = header["Number"]
    for row in range(sheet.max_row, 1, -1):
        building_id = str(sheet.cell(row, number_column).value).strip()
        if building_id not in keep_ids:
            sheet.delete_rows(row, 1)

    audit = workbook.create_sheet("Exclusion_Audit")
    audit.append(["Item", "Value"])
    audit.append(["Rule", "Remove every building belonging to the 61 coordinate-overlapping campuses"])
    audit.append(["Excluded campuses", len(excluded_campuses)])
    audit.append(["Remaining buildings", len(keep_ids)])
    audit.append(["Remaining campuses", len({campus_id(value) for value in keep_ids})])
    audit.append([])
    audit.append(["Excluded campus ID"])
    for value in sorted(excluded_campuses):
        audit.append([value])
    workbook.save(destination)


def filter_year_workbook(keep_ids: set[str]) -> int:
    if not SOURCE_YEAR.exists():
        return 0
    destination = OUT / "source-data" / SOURCE_YEAR.name
    workbook = load_workbook(SOURCE_YEAR)
    sheet = workbook.worksheets[0]
    # The first column is the building identifier in the source simulation-input workbook.
    removed = 0
    for row in range(sheet.max_row, 1, -1):
        building_id = str(sheet.cell(row, 1).value).strip()
        if building_id and building_id not in keep_ids:
            sheet.delete_rows(row, 1)
            removed += 1
    workbook.save(destination)
    return removed


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    overlap = pd.read_csv(OVERLAP)
    overlap["excluded_campus_id"] = overlap["training_id"].astype(str).map(campus_id)
    excluded_campuses = set(overlap["excluded_campus_id"].unique())
    if len(excluded_campuses) != 61:
        raise ValueError(f"Expected 61 unique campuses, found {len(excluded_campuses)}")

    numerical = pd.read_excel(SOURCE_XLSX, sheet_name="selected_variables")
    building_ids = numerical["Number"].astype(str).str.strip()
    campuses = building_ids.map(campus_id)
    keep_mask = ~campuses.isin(excluded_campuses)
    keep_ids = set(building_ids[keep_mask])
    removed_ids = set(building_ids[~keep_mask])
    remaining_campuses = {campus_id(value) for value in keep_ids}

    filter_workbook(keep_ids, excluded_campuses)
    filter_year_workbook(keep_ids)

    voxels = np.load(SOURCE_VOXELS, mmap_mode="r")
    if len(voxels) != len(numerical):
        raise ValueError(f"Voxel rows {len(voxels)} do not match numerical rows {len(numerical)}")
    (OUT / "geometry").mkdir(parents=True, exist_ok=True)
    np.save(OUT / "geometry" / "training_voxels_32.npy", np.asarray(voxels[keep_mask.to_numpy()]))

    with SOURCE_BUILDINGS.open("r", encoding="utf-8") as handle:
        buildings = json.load(handle)
    filtered_buildings = [record for record in buildings if str(record.get("id")) in keep_ids]
    (OUT / "metadata").mkdir(parents=True, exist_ok=True)
    with (OUT / "metadata" / "buildings.json").open("w", encoding="utf-8") as handle:
        json.dump(filtered_buildings, handle, ensure_ascii=False, separators=(",", ":"))

    gh = pd.read_csv(SOURCE_GH, encoding="utf-8-sig")
    gh_id_column = "obj_object_name_prefix"
    gh_filtered = gh[gh[gh_id_column].astype(str).isin(keep_ids)].copy()
    gh_filtered.to_csv(OUT / "metadata" / "gh_building_energy_inputs.csv", index=False, encoding="utf-8-sig")

    ts_out = OUT / "time-series" / "Time series data"
    ts_out.mkdir(parents=True, exist_ok=True)
    ts_missing = []
    ts_modes = {"hardlink": 0, "copy": 0, "existing": 0}
    for building_id in sorted(keep_ids):
        source = SOURCE_TS / f"{building_id}.xlsx"
        if not source.exists():
            source = SOURCE_TS_FALLBACK / f"{building_id}.xlsx"
        if not source.exists():
            ts_missing.append(building_id)
            continue
        mode = hardlink_or_copy(source, ts_out / source.name)
        ts_modes[mode] += 1

    obj_out = OUT / "geometry" / "per_campus_obj"
    obj_out.mkdir(parents=True, exist_ok=True)
    obj_missing = []
    obj_modes = {"hardlink": 0, "copy": 0, "existing": 0}
    for campus in sorted(remaining_campuses):
        source = SOURCE_OBJ / f"{campus}.obj"
        if not source.exists():
            obj_missing.append(campus)
            continue
        mode = hardlink_or_copy(source, obj_out / source.name)
        obj_modes[mode] += 1

    overlap.to_csv(OUT / "excluded_61_campuses.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(
        {
            "building_id": building_ids,
            "campus_id": campuses,
            "status": np.where(keep_mask, "retained", "excluded_external_overlap"),
        }
    ).to_csv(OUT / "building_exclusion_manifest.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame({"missing_time_series_building_id": ts_missing}).to_csv(
        OUT / "missing_time_series.csv", index=False, encoding="utf-8-sig"
    )
    pd.DataFrame({"missing_campus_obj": obj_missing}).to_csv(
        OUT / "missing_campus_obj.csv", index=False, encoding="utf-8-sig"
    )

    summary = {
        "source_buildings": int(len(numerical)),
        "source_campuses": int(campuses.nunique()),
        "excluded_external_campuses": int(len(excluded_campuses)),
        "excluded_buildings": int((~keep_mask).sum()),
        "remaining_campuses": int(len(remaining_campuses)),
        "remaining_buildings": int(keep_mask.sum()),
        "remaining_voxel_rows": int(keep_mask.sum()),
        "remaining_time_series_files": int(sum(ts_modes.values())),
        "remaining_campus_obj_files": int(sum(obj_modes.values())),
        "missing_time_series_files": int(len(ts_missing)),
        "missing_campus_obj_files": int(len(obj_missing)),
        "time_series_file_modes": ts_modes,
        "campus_obj_file_modes": obj_modes,
        "original_files_modified": False,
    }
    (OUT / "dataset_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    pd.DataFrame([summary | {"time_series_file_modes": str(ts_modes), "campus_obj_file_modes": str(obj_modes)}]).to_csv(
        OUT / "dataset_summary.csv", index=False, encoding="utf-8-sig"
    )

    # Final leakage and alignment assertions.
    remaining_sheet = pd.read_excel(OUT / "source-data" / "numerical data.xlsx", sheet_name="selected_variables")
    remaining_ids = remaining_sheet["Number"].astype(str).str.strip()
    assert len(remaining_ids) == keep_mask.sum()
    assert not remaining_ids.map(campus_id).isin(excluded_campuses).any()
    assert np.load(OUT / "geometry" / "training_voxels_32.npy", mmap_mode="r").shape[0] == len(remaining_ids)
    assert not removed_ids.intersection(set(remaining_ids))
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
