from __future__ import annotations

from concurrent.futures import ProcessPoolExecutor
from pathlib import Path

import numpy as np
import pandas as pd
from numpy.lib.format import open_memmap
from openpyxl import load_workbook

ROOT = Path(r"C:\Users\DELL\Documents\论文\internal_two_route_rebuild")
TABLE = Path(r"C:\Users\DELL\Desktop\SCI8\数据集完成\numerical data.xlsx")
TS = Path(r"C:\Users\DELL\Desktop\SCI8\数据集完成\Time series data")
TS_FALLBACK = Path(r"C:\Users\DELL\Desktop\SCI8\Time series data")
OUT = ROOT / "route2_cache"


def read_one(item):
    i, building_id = item
    path = TS / f"{building_id}.xlsx"
    if not path.exists():
        path = TS_FALLBACK / f"{building_id}.xlsx"
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(min_row=3, max_row=8762, min_col=3, max_col=16, values_only=True))
    wb.close()
    raw = np.asarray(rows, dtype=np.float32)
    if raw.shape != (8760, 14):
        return i, None, None
    day, hour = raw[:, 0], raw[:, 1]
    weather = np.column_stack([
        np.sin(2*np.pi*day/365), np.cos(2*np.pi*day/365),
        np.sin(2*np.pi*hour/24), np.cos(2*np.pi*hour/24), raw[:, 2:8],
    ]).astype(np.float32)
    target = raw[:, [8, 9, 13]].astype(np.float32)
    return i, weather, target


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    df = pd.read_excel(TABLE, sheet_name="selected_variables")
    ids = df["Number"].astype(str).tolist(); n = len(ids)
    weather = open_memmap(OUT/"weather.npy", mode="w+", dtype="float32", shape=(n,8760,10))
    targets = open_memmap(OUT/"targets.npy", mode="w+", dtype="float32", shape=(n,8760,3))
    weather[:]=np.nan; targets[:]=np.nan; valid=np.zeros(n,dtype=bool)
    with ProcessPoolExecutor(max_workers=10) as pool:
        for done,(i,w,y) in enumerate(pool.map(read_one, enumerate(ids), chunksize=1),1):
            if w is not None:
                weather[i]=w; targets[i]=y; valid[i]=True
            if done==1 or done%100==0 or done==n:
                weather.flush(); targets.flush(); print(f"cached {done}/{n}",flush=True)
    pd.DataFrame({"row_index":np.arange(n),"building_id":ids}).to_csv(OUT/"index.csv",index=False)
    np.save(OUT/"valid.npy",valid)


if __name__=="__main__": main()
